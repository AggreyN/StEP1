"""A resume (schemas.ResumeDoc) as a file: PDF or DOCX.

Both are single-column and plain on purpose, so applicant tracking systems
read them in order: name, contact line, then each section's entries. An
entry with a heading is an experience-style block (heading with dates on the
right, then sub with place on the right, then its lines as bullets); an
entry without a heading is plain rows (a skills section). **bold** in a line
is bold.

The PDF is ported from the owner's resume skill: Helvetica, letter size,
stepping the font and margins down a fixed ladder until the resume fits one
page. If even the smallest step overflows, the smallest is returned anyway
(two pages beat none) and `pages` says so.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

# (body font size, side/top margin in inches, line leading multiplier)
FIT_LADDER = [
    (10.5, 0.6, 1.22),
    (10.25, 0.55, 1.2),
    (10.0, 0.5, 1.18),
    (9.75, 0.5, 1.16),
    (9.5, 0.45, 1.15),
]


@dataclass
class Rendered:
    data: bytes
    pages: int
    font_size: float


def _md(text: str) -> str:
    """Escape for ReportLab's markup, then **bold** -> <b>."""
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)


def _pdf_once(doc: dict, font: float, margin: float, lead: float) -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=margin * inch,
        rightMargin=margin * inch,
        topMargin=margin * 0.8 * inch,
        bottomMargin=margin * 0.8 * inch,
        title=f"{doc['name']} Resume",
        author=doc["name"],
        creator="StEP1",
        producer="StEP1",
    )
    width = letter[0] - 2 * margin * inch
    leading = font * lead
    style = {
        "name": ParagraphStyle(
            "name", fontName="Helvetica-Bold", fontSize=font + 8, leading=font + 11,
            alignment=TA_CENTER,
        ),
        "contact": ParagraphStyle(
            "contact", fontName="Helvetica", fontSize=font - 1, leading=leading,
            alignment=TA_CENTER,
        ),
        "section": ParagraphStyle(
            "section", fontName="Helvetica-Bold", fontSize=font + 0.5, leading=leading,
            spaceBefore=font * 0.55,
        ),
        "head": ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=font, leading=leading),
        "sub": ParagraphStyle("sub", fontName="Helvetica-Oblique", fontSize=font, leading=leading),
        "right": ParagraphStyle(
            "right", fontName="Helvetica", fontSize=font, leading=leading, alignment=TA_RIGHT
        ),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=font, leading=leading),
        "bullet": ParagraphStyle(
            "bullet", fontName="Helvetica", fontSize=font, leading=leading, leftIndent=12,
            bulletIndent=3,
        ),
    }  # fmt: skip
    story = [Paragraph(_md(doc["name"]), style["name"])]
    if doc.get("contact"):
        story.append(Paragraph(_md("  |  ".join(doc["contact"])), style["contact"]))

    def row(left: str, right: str, left_style):
        if not right:
            return Paragraph(_md(left), left_style)
        table = Table(
            [[Paragraph(_md(left), left_style), Paragraph(_md(right), style["right"])]],
            colWidths=[width * 0.74, width * 0.26],
        )
        table.setStyle(
            TableStyle(
                [
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        return table

    for section in doc.get("sections", []):
        story.append(Paragraph(_md(section["title"].upper()), style["section"]))
        story.append(
            HRFlowable(
                width="100%", thickness=0.6, color="black", spaceBefore=1, spaceAfter=font * 0.25
            )
        )
        for entry in section.get("entries", []):
            block = []
            if entry.get("heading"):
                block.append(row(entry["heading"], entry.get("right", ""), style["head"]))
            if entry.get("sub"):
                block.append(row(entry["sub"], entry.get("sub_right", ""), style["sub"]))
            for line in entry.get("lines", []):
                if entry.get("heading"):
                    block.append(Paragraph(_md(line), style["bullet"], bulletText="•"))
                else:
                    block.append(Paragraph(_md(line), style["body"]))
            block.append(Spacer(1, font * 0.3))
            story.append(KeepTogether(block))
    pdf.build(story)
    return buf.getvalue()


def pdf(doc: dict) -> Rendered:
    """One page if it can be done on the ladder; otherwise the smallest step."""
    from pypdf import PdfReader

    data, pages, font = b"", 0, 0.0
    for font, margin, lead in FIT_LADDER:
        data = _pdf_once(doc, font, margin, lead)
        pages = len(PdfReader(io.BytesIO(data)).pages)
        if pages == 1:
            break
    return Rendered(data=data, pages=pages, font_size=font)


def docx(doc: dict) -> Rendered:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.6)
    section.top_margin = section.bottom_margin = Inches(0.5)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    usable = section.page_width - section.left_margin - section.right_margin
    document.core_properties.title = f"{doc['name']} Resume"
    document.core_properties.author = doc["name"]

    def runs(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
        for n, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                run = paragraph.add_run(part)
                run.bold = bold or n % 2 == 1
                run.italic = italic

    def row(left: str, right: str, bold: bool = False, italic: bool = False) -> None:
        p = document.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
        runs(p, left, bold=bold, italic=italic)
        if right:
            p.add_run("\t")
            runs(p, right, italic=italic)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(doc["name"])
    run.bold, run.font.size = True, Pt(18)
    if doc.get("contact"):
        contact = document.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(contact, "  |  ".join(doc["contact"]))

    for block in doc.get("sections", []):
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        title = heading.add_run(block["title"].upper())
        title.bold, title.font.size = True, Pt(11)
        for entry in block.get("entries", []):
            if entry.get("heading"):
                row(entry["heading"], entry.get("right", ""), bold=True)
            if entry.get("sub"):
                row(entry["sub"], entry.get("sub_right", ""), italic=True)
            for line in entry.get("lines", []):
                style = "List Bullet" if entry.get("heading") else None
                runs(document.add_paragraph(style=style), line)
    buf = io.BytesIO()
    document.save(buf)
    return Rendered(data=buf.getvalue(), pages=0, font_size=10.5)
